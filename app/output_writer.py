"""
Stage 5b – Output Writer.
  - catalog_clean.xlsx : validated records (thumbnail + hyperlinked image paths) and an
                         Exceptions sheet for reviewers.
  - catalog_clean.json : validated records, machine-readable.
  - exceptions.json    : quarantined products, invalid VLM records, page errors, skipped pages.
  - exceptions.log     : the same, human-readable, with exact failure reasons.
Image paths are written relative to the job output directory so the folder stays portable.
"""
from __future__ import annotations

import io
import json
import logging
from datetime import datetime, timezone
from pathlib import Path
from typing import List

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image

from app.schemas import JobStatus, PageIssue, ProcessedProduct

logger = logging.getLogger(__name__)

_HEADER_FILL = PatternFill("solid", fgColor="1E3A5F")
_HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
_ALT_FILL = PatternFill("solid", fgColor="EEF2F7")
_LINK_FONT = Font(color="1155CC", underline="single")
_THUMB_PX = 72


def _slabs_text(pp: ProcessedProduct) -> str:
    return "; ".join(
        f"{s.min_qty}{'-' + str(s.max_qty) if s.max_qty else '+'} pcs: {s.price:g}"
        for s in sorted(pp.extracted.price_slabs, key=lambda s: s.min_qty)
    )


def _changes_text(pp: ProcessedProduct) -> str:
    parts = []
    for field, c in pp.changes.items():
        delta = f" (Δ {c['delta']:+g})" if "delta" in c else ""
        parts.append(f"{field}: {c['old']} → {c['new']}{delta}")
    return "; ".join(parts)


def product_record(pp: ProcessedProduct) -> dict:
    e = pp.extracted
    return {
        "page": pp.page_number,
        "source": pp.source,
        "tier": pp.tier,
        "product_name": e.product_name,
        "sku_or_model": e.sku_or_model,
        "brand": e.brand,
        "category": e.category,
        "description": e.description,
        "mrp": e.mrp,
        "b2b_price": e.b2b_price,
        "moq": e.moq,
        "color_variants": e.color_variants,
        "specifications": e.specifications.model_dump(),
        "price_slabs": [s.model_dump() for s in e.price_slabs],
        "images": [i.model_dump() for i in e.images],
        "image_paths": pp.image_paths,
        "original_image_paths": pp.original_image_paths,
        "background_removed": pp.background_removed,
        "image_warnings": pp.image_warnings,
        "raw_text_snippet": e.raw_text_snippet,
        "confidence_score": pp.validation.confidence_score,
        "check_scores": pp.validation.check_scores,
        "text_source": pp.validation.text_source,
        "warnings" if pp.validation.passed else "failure_reasons": pp.validation.failure_reasons,
        "change_status": pp.change_status,
        "changes": pp.changes,
    }


# ── Excel ─────────────────────────────────────────────────────────────────────

_CLEAN_COLUMNS = [
    # (header, width)
    ("Preview", 12), ("Page", 7), ("Product Name", 30), ("SKU / Model", 16), ("Brand", 14), ("Category", 16),
    ("Description", 40), ("MRP", 10), ("B2B Price", 10), ("MOQ", 8), ("Price Slabs", 24),
    ("Color Variants", 18), ("Material", 16), ("Capacity / Size", 14), ("Dimensions", 16),
    ("Confidence", 11), ("Change Status", 12), ("Changes", 30), ("Primary Image", 32),
    ("All Images", 40), ("Background Removed", 12), ("Warnings", 40),
]


_PRIMARY_IMAGE_COL = [h for h, _ in _CLEAN_COLUMNS].index("Primary Image") + 1


def _write_header(ws, columns) -> None:
    for idx, (header, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=idx, value=header)
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "B2"


def _thumbnail(path: Path) -> XLImage | None:
    try:
        with Image.open(path) as img:
            img = img.convert("RGBA")
            img.thumbnail((_THUMB_PX, _THUMB_PX))
            buf = io.BytesIO()
            img.save(buf, format="PNG")
        buf.seek(0)
        return XLImage(buf)
    except Exception:
        return None


def export_excel(
    validated: List[ProcessedProduct],
    exceptions: List[ProcessedProduct],
    output_path: Path,
) -> Path:
    job_dir = output_path.parent
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Validated Catalog"
    _write_header(ws, _CLEAN_COLUMNS)

    for r, pp in enumerate(validated, 2):
        e, s = pp.extracted, pp.extracted.specifications
        values = [
            None, pp.page_number, e.product_name, e.sku_or_model, e.brand, e.category, e.description,
            e.mrp, e.b2b_price, e.moq, _slabs_text(pp), ", ".join(e.color_variants),
            s.material, s.capacity_or_size, s.dimensions, pp.validation.confidence_score,
            pp.change_status, _changes_text(pp), None, "\n".join(pp.image_paths),
            {True: "Yes", False: "No", None: "—"}[pp.background_removed],
            "\n".join(pp.validation.failure_reasons + pp.image_warnings),
        ]
        fill = _ALT_FILL if r % 2 == 0 else None
        for c, v in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if fill:
                cell.fill = fill

        if pp.image_paths:
            link = ws.cell(row=r, column=_PRIMARY_IMAGE_COL, value=pp.image_paths[0])
            link.hyperlink = pp.image_paths[0]  # relative to the workbook's folder
            link.font = _LINK_FONT
            thumb = _thumbnail(job_dir / pp.image_paths[0])
            if thumb is not None:
                ws.add_image(thumb, f"A{r}")
                ws.row_dimensions[r].height = 60

    # Exceptions sheet for reviewers
    ex = wb.create_sheet("Exceptions")
    ex_cols = [("Page", 7), ("Product Name", 30), ("SKU / Model", 16), ("MRP", 10), ("B2B Price", 10),
               ("MOQ", 8), ("Confidence", 11), ("Failure Reasons", 70), ("Raw Text Snippet", 50)]
    _write_header(ex, ex_cols)
    for r, pp in enumerate(exceptions, 2):
        e = pp.extracted
        for c, v in enumerate([pp.page_number, e.product_name, e.sku_or_model, e.mrp, e.b2b_price, e.moq,
                               pp.validation.confidence_score, "\n".join(pp.validation.failure_reasons),
                               e.raw_text_snippet], 1):
            ex.cell(row=r, column=c, value=v).alignment = Alignment(wrap_text=True, vertical="top")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    logger.info("Excel exported: %s (%d validated, %d exceptions)", output_path, len(validated), len(exceptions))
    return output_path


# ── JSON ──────────────────────────────────────────────────────────────────────

def export_clean_json(validated: List[ProcessedProduct], job: JobStatus, output_path: Path) -> Path:
    payload = {
        "job_id": job.job_id,
        "source": job.filename,
        "generated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "model_tiers": job.tiers,
        "cost_usd": job.cost_usd,
        "products": [product_record(p) for p in validated],
    }
    output_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False), encoding="utf-8")
    return output_path


def export_exceptions(
    exceptions: List[ProcessedProduct],
    invalid_records: List[dict],
    job: JobStatus,
    output_path: Path,
) -> Path:
    payload = {
        "job_id": job.job_id,
        "source": job.filename,
        "products": [product_record(p) for p in exceptions],
        "invalid_vlm_records": invalid_records,
        "page_errors": [i.model_dump() for i in job.page_errors],
        "skipped_pages": [i.model_dump() for i in job.skipped_pages],
    }
    output_path.write_text(json.dumps(payload, indent=2, ensure_ascii=False, default=str), encoding="utf-8")
    logger.info("Exceptions written: %s (%d products)", output_path, len(exceptions))
    return output_path


def export_exceptions_log(
    exceptions: List[ProcessedProduct],
    validated: List[ProcessedProduct],
    invalid_records: List[dict],
    job: JobStatus,
    output_path: Path,
) -> Path:
    lines = [
        f"Catalog Ingestion Engine exceptions log",
        f"Job:       {job.job_id}",
        f"Source:    {job.filename}",
        f"Models:    {job.tiers}",
        f"AI cost:   ${job.cost_usd:.4f} (₹{job.cost_inr:.2f})",
        f"Generated: {datetime.now(timezone.utc).isoformat(timespec='seconds')}",
        f"Summary:   {len(validated)} validated, {len(exceptions)} quarantined, "
        f"{len(job.page_errors)} page error(s), {len(job.skipped_pages)} skipped page(s)",
        "",
    ]

    def section(title: str, issues: List[PageIssue]) -> None:
        if issues:
            lines.append(f"== {title} ==")
            lines.extend(f"[page {i.page_number}] {i.source}: {i.reason}" for i in issues)
            lines.append("")

    section("PAGE ERRORS (no products extracted — re-run or review manually)", job.page_errors)
    section("PAGE WARNINGS (review: possibly missed products)", job.page_warnings)

    escalated = [p for p in job.pages if len(p.attempts) > 1]
    if escalated:
        lines.append("== MODEL ESCALATIONS ==")
        for p in escalated:
            lines.append(f"[page {p.page_number}] " + " | ".join(f"{a.tier}: {a.outcome}" for a in p.attempts))
        lines.append("")

    if exceptions:
        lines.append("== QUARANTINED PRODUCTS ==")
        for pp in sorted(exceptions, key=lambda p: p.page_number):
            e = pp.extracted
            lines.append(f"[page {pp.page_number}] {e.product_name} (SKU: {e.sku_or_model or '—'}) "
                         f"confidence={pp.validation.confidence_score:.2f}")
            lines.append(f"    mrp={e.mrp} b2b_price={e.b2b_price} moq={e.moq}")
            lines.extend(f"    - {r}" for r in pp.validation.failure_reasons)
            if e.raw_text_snippet:
                lines.append(f"    snippet: {e.raw_text_snippet!r}")
        lines.append("")

    if invalid_records:
        lines.append("== MALFORMED VLM RECORDS ==")
        for rec in invalid_records:
            lines.append(f"[page {rec.get('page')}] {rec.get('error')}")
        lines.append("")

    warned = [p for p in validated if p.image_warnings]
    if warned:
        lines.append("== IMAGE WARNINGS (validated products) ==")
        for pp in warned:
            lines.extend(f"[page {pp.page_number}] {pp.extracted.product_name}: {w}" for w in pp.image_warnings)
        lines.append("")

    section("SKIPPED PAGES", job.skipped_pages)

    output_path.write_text("\n".join(lines), encoding="utf-8")
    return output_path
