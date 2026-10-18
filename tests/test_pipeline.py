"""End-to-end pipeline test on a generated PDF with a fake VLM (no API calls)."""
import asyncio
import io
import json

import openpyxl
import pymupdf as fitz
import pytest
from PIL import Image, ImageDraw

from app.pdf_renderer import enumerate_pages, extract_page_text, parse_page_range, skip_reason_for
from app.pipeline import PipelineOptions, run_pipeline
from app.schemas import PageExtractionResult
from app.vlm_extractor import BaseExtractor, ExtractionError, parse_page_payload


def _product_png() -> bytes:
    """A red mug on a white background."""
    img = Image.new("RGB", (400, 400), "white")
    d = ImageDraw.Draw(img)
    d.rectangle([110, 90, 290, 320], fill=(200, 30, 30))
    d.ellipse([270, 150, 350, 250], outline=(200, 30, 30), width=18)
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def catalog_pdf(tmp_path):
    doc = fitz.open()
    cover = doc.new_page(width=595, height=842)
    cover.insert_text((200, 400), "GIFT CATALOGUE 2026", fontsize=24)

    page = doc.new_page(width=595, height=842)
    page.insert_image(fitz.Rect(60, 80, 300, 320), stream=_product_png())
    page.insert_text((60, 360), "Red Ceramic Mug RM-330", fontsize=14)
    page.insert_text((60, 385), "Capacity 330 ml, stoneware", fontsize=11)
    page.insert_text((60, 410), "MRP Rs. 1,450/-   Corporate price 899   MOQ 50 pcs", fontsize=11)
    page.insert_text((320, 360), "Blue Steel Bottle BB-750", fontsize=14)
    page.insert_text((320, 410), "MRP Rs. 999", fontsize=11)

    path = tmp_path / "catalog.pdf"
    doc.save(str(path))
    return path


class FakeExtractor(BaseExtractor):
    provider = "fake"

    def __init__(self):
        super().__init__("fake-vlm")
        self.pages_seen = []

    async def extract_page(self, image_bytes, page_number, native_text="", text_source="native"):
        self.pages_seen.append(page_number)
        # Image box: PDF rect (60,80)-(300,320) on a 595x842 page, as 0-1000 [ymin,xmin,ymax,xmax].
        box = [80 / 842 * 1000, 60 / 595 * 1000, 320 / 842 * 1000, 300 / 595 * 1000]
        return parse_page_payload({"products": [
            {   # fully correct → validated
                "product_name": "Red Ceramic Mug", "sku_or_model": "RM-330", "category": "Drinkware",
                "description": None, "mrp": 1450, "b2b_price": 899, "moq": 50, "color_variants": ["Red"],
                "specifications": {"material": "stoneware", "capacity_or_size": "330 ml", "dimensions": None},
                "images": [{"box_2d": box, "type": "primary"}], "price_slabs": [],
                "raw_text_snippet": "Red Ceramic Mug RM-330",
            },
            {   # hallucinated b2b price → quarantined
                "product_name": "Blue Steel Bottle", "sku_or_model": "BB-750", "mrp": 999, "b2b_price": 799,
                "images": [], "raw_text_snippet": "Blue Steel Bottle BB-750",
            },
            {"product_name": None, "images": "broken"},  # malformed → invalid record
        ]}, page_number)


class FailingExtractor(BaseExtractor):
    provider = "fake"

    async def extract_page(self, *a, **kw) -> PageExtractionResult:
        raise ExtractionError("boom")


def _run(pdf, tmp_path, *extractors, **opts):
    options = PipelineOptions(output_root=tmp_path / "out", **{"max_concurrent": 2, **opts})
    tiers = [e for e in extractors if e is not None] or None
    return asyncio.run(run_pipeline(pdf, options, job_id="job1", tiers=tiers))


def test_end_to_end(catalog_pdf, tmp_path):
    fake = FakeExtractor()
    job = _run(catalog_pdf, tmp_path, fake)

    assert job.status == "done", job.error_message
    assert fake.pages_seen == [2]                       # cover page skipped before any VLM call
    assert [p.page_number for p in job.skipped_pages] == [1]
    assert (job.validated_products, job.exception_products) == (1, 1)
    assert job.new_products == 1

    out = tmp_path / "out" / "job1"
    clean = json.loads((out / "catalog_clean.json").read_text(encoding="utf-8"))["products"]
    assert clean[0]["sku_or_model"] == "RM-330" and clean[0]["b2b_price"] == 899
    img_path = out / clean[0]["image_paths"][0]
    assert img_path.name == "RM-330_0.png"
    with Image.open(img_path) as img:
        assert img.mode == "RGBA"
        assert max(img.size) > 600                  # ~250 px native photo upscaled (capped at 3×)
        alpha = img.getchannel("A")
        assert alpha.getextrema()[0] == 0           # some background removed
        assert clean[0]["background_removed"] is True
    assert any("low resolution" in w for w in clean[0]["image_warnings"])
    assert (out / clean[0]["original_image_paths"][0]).exists()     # original crop kept for review

    exc = json.loads((out / "exceptions.json").read_text(encoding="utf-8"))
    assert exc["products"][0]["sku_or_model"] == "BB-750"
    assert any("799" in r for r in exc["products"][0]["failure_reasons"])
    assert len(exc["invalid_vlm_records"]) == 1
    assert "BB-750" in (out / "exceptions.log").read_text(encoding="utf-8")

    wb = openpyxl.load_workbook(out / "catalog_clean.xlsx")
    ws = wb["Validated Catalog"]
    assert ws.cell(row=2, column=4).value == "RM-330"
    from app.output_writer import _PRIMARY_IMAGE_COL
    assert ws.cell(row=2, column=_PRIMARY_IMAGE_COL).hyperlink.target == "images/RM-330_0.png"
    assert wb["Exceptions"].cell(row=2, column=3).value == "BB-750"


def test_rerun_detects_unchanged_then_price_update(catalog_pdf, tmp_path):
    _run(catalog_pdf, tmp_path, FakeExtractor())
    job = _run(catalog_pdf, tmp_path, FakeExtractor())
    assert (job.new_products, job.unchanged_products) == (0, 1)

    class Repriced(FakeExtractor):
        async def extract_page(self, *a, **kw):
            res = await super().extract_page(*a, **kw)
            res.products[0].b2b_price = 1450.0  # also printed on the page, so it verifies
            return res

    job = _run(catalog_pdf, tmp_path, Repriced())
    assert job.updated_products == 1
    master = json.loads((tmp_path / "out" / "master_catalog.json").read_text(encoding="utf-8"))
    hist = master["products"]["rm-330"]["history"][-1]["changes"]["b2b_price"]
    assert hist == {"old": 899, "new": 1450.0, "delta": 551.0}


def test_duplicates_conflict_only_on_commercial_terms():
    from app.catalog_store import resolve_duplicates
    from app.schemas import ExtractedProduct, ProcessedProduct, ValidationResult

    def pp(page, **kw):
        e = ExtractedProduct(product_name="Ellipse Clock", mrp=3499, **kw)
        return ProcessedProduct(page_number=page, extracted=e,
                                validation=ValidationResult(confidence_score=1, passed=True))

    kept, quarantined = resolve_duplicates([pp(27, category="Alarm Clock"), pp(28, category="Clock")])
    assert len(kept) == 1 and not quarantined                  # wording differs, prices agree
    kept, quarantined = resolve_duplicates([pp(27, b2b_price=2999), pp(28, b2b_price=2799)])
    assert not kept and len(quarantined) == 2                  # price conflict


def test_fatal_error_stops_remaining_pages(catalog_pdf, tmp_path):
    from app.vlm_extractor import FatalExtractionError

    class NoCredits(BaseExtractor):
        provider = "fake"
        calls = 0

        async def extract_page(self, *a, **kw):
            NoCredits.calls += 1
            raise FatalExtractionError("out of API credits")

    job = _run(catalog_pdf, tmp_path, NoCredits("x"), skip_non_product=False, max_concurrent=1)
    assert NoCredits.calls == 1                               # second page never hits the API
    assert job.status == "error" and len(job.page_errors) == 2
    assert "all model tiers disabled" in job.page_errors[1].reason


def test_all_pages_failing_marks_job_error(catalog_pdf, tmp_path):
    job = _run(catalog_pdf, tmp_path, FailingExtractor("x"), skip_non_product=False)
    assert job.status == "error"
    assert len(job.page_errors) == 2


def test_missing_api_keys_is_a_clear_error(catalog_pdf, tmp_path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "anthropic_api_key", "")
    monkeypatch.setattr(settings, "gemini_api_key", "")
    job = _run(catalog_pdf, tmp_path, None)
    assert job.status == "error" and "No usable model tiers" in job.error_message


# ── cost ladder ──────────────────────────────────────────────────────────────

class HallucinatingCheap(FakeExtractor):
    """Cheap tier that gets the mug's B2B price wrong (899 → 989)."""

    def __init__(self, model="cheap-model"):
        super().__init__()
        self.model = model

    async def extract_page(self, *a, **kw):
        res = await super().extract_page(*a, **kw)
        res.products[0].b2b_price = 989.0
        self.input_tokens += 1000
        self.output_tokens += 100
        return res


def test_failed_verification_escalates_to_next_tier(catalog_pdf, tmp_path):
    cheap, strong = HallucinatingCheap(), FakeExtractor()
    job = _run(catalog_pdf, tmp_path, cheap, strong)
    assert job.status == "done"
    assert cheap.pages_seen == [2] and strong.pages_seen == [2]
    page = job.pages[0]
    assert page.tier_used == "fake:fake-vlm"
    assert page.attempts[0].outcome.startswith("escalated:")
    clean = json.loads((tmp_path / "out" / "job1" / "catalog_clean.json").read_text(encoding="utf-8"))
    assert clean["products"][0]["b2b_price"] == 899 and clean["products"][0]["tier"] == "fake:fake-vlm"
    assert [t.pages_attempted for t in job.tier_stats] == [1, 1]
    assert (tmp_path / "out" / "job1" / "pages" / "p2.jpg").exists()
    assert (tmp_path / "out" / "job1" / "job.json").exists()


def test_agreeing_tiers_stop_escalating(catalog_pdf, tmp_path):
    a, b, premium = HallucinatingCheap("a"), HallucinatingCheap("b"), FakeExtractor()
    job = _run(catalog_pdf, tmp_path, a, b, premium)
    assert premium.pages_seen == []                       # same answer twice → premium not called
    assert "same result as previous tier" in job.pages[0].attempts[-1].outcome
    assert job.validated_products == 0                    # still quarantined


class UnreadableName(FakeExtractor):
    """Returns a name whose first word is not in the page text (as with a stylised logo OCR can't read)."""

    def __init__(self, model):
        super().__init__()
        self.model = model

    async def extract_page(self, *a, **kw):
        res = await super().extract_page(*a, **kw)
        res.products[0].product_name = "ZYRO Red Ceramic Mug"
        res.products[1].b2b_price = None   # bottle correct, so only the name question remains
        res.invalid_records = []
        return res


def test_two_models_agreeing_on_unreadable_name_are_accepted_with_warning(catalog_pdf, tmp_path):
    a, b = UnreadableName("a"), UnreadableName("b")
    job = _run(catalog_pdf, tmp_path, a, b)
    assert job.validated_products == 2
    assert "accepted by agreement" in job.pages[0].attempts[-1].outcome
    clean = json.loads((tmp_path / "out" / "job1" / "catalog_clean.json").read_text(encoding="utf-8"))
    mug = next(p for p in clean["products"] if p["product_name"].startswith("ZYRO"))
    assert any("confirmed by agreement" in w for w in mug["warnings"])


def test_name_consensus_can_be_disabled(catalog_pdf, tmp_path, monkeypatch):
    from app.config import settings
    monkeypatch.setattr(settings, "name_consensus", False)
    job = _run(catalog_pdf, tmp_path, UnreadableName("a"), UnreadableName("b"))
    assert job.validated_products == 1      # bottle only; the unconfirmed name stays quarantined


def test_consensus_never_overrides_a_price_failure(catalog_pdf, tmp_path):
    job = _run(catalog_pdf, tmp_path, HallucinatingCheap("a"), HallucinatingCheap("b"))
    assert job.validated_products == 0


def test_blocked_escalation_is_reported_as_rerun_not_as_accuracy(catalog_pdf, tmp_path):
    from app.vlm_extractor import FatalExtractionError

    class OutOfQuota(BaseExtractor):
        provider = "fake"

        async def extract_page(self, *a, **kw):
            raise FatalExtractionError("quota exhausted")

    job = _run(catalog_pdf, tmp_path, HallucinatingCheap(), OutOfQuota("strong"))
    assert job.pages[0].needs_rerun
    assert any(w.reason.startswith("needs re-run") for w in job.page_warnings)

    # A ladder that simply ends (no stronger tier configured) is a genuine verdict, not a re-run.
    job = _run(catalog_pdf, tmp_path, HallucinatingCheap())
    assert not job.pages[0].needs_rerun


def test_job_cost_cap_stops_further_model_calls(catalog_pdf, tmp_path, monkeypatch):
    from app import config
    monkeypatch.setitem(config.MODEL_PRICING, "cheap-model", (1000.0, 1000.0))  # 1,100 tokens ≈ $1.10
    cheap, strong = HallucinatingCheap(), FakeExtractor()
    job = _run(catalog_pdf, tmp_path, cheap, strong, max_cost_usd=0.5)
    assert strong.pages_seen == []                      # escalation refused: cap already reached
    assert "cost cap" in job.pages[0].attempts[-1].outcome
    assert job.pages[0].needs_rerun


def test_unavailable_tier_falls_through(catalog_pdf, tmp_path):
    from app.vlm_extractor import FatalExtractionError

    class Down(BaseExtractor):
        provider = "fake"

        async def extract_page(self, *a, **kw):
            raise FatalExtractionError("model retired")

    job = _run(catalog_pdf, tmp_path, Down("down"), FakeExtractor())
    assert job.status == "done" and job.validated_products == 1
    assert job.tier_stats[0].disabled_reason == "model retired"


# ── Stage 1 helpers ──────────────────────────────────────────────────────────

def test_folder_input_and_page_range(tmp_path, catalog_pdf):
    folder = tmp_path / "pages"
    folder.mkdir()
    (folder / "page10.png").write_bytes(_product_png())
    (folder / "page2.png").write_bytes(_product_png())
    (folder / "page2 (1).png").write_bytes(_product_png())                # byte-identical copy
    (folder / "page10.png").write_bytes(_product_png() + b"\0")
    refs = enumerate_pages(folder)
    assert [r.source.name for r in refs] == ["page2.png", "page10.png"]   # natural sort, copy dropped
    assert extract_page_text(refs[0], use_ocr=False).text_source == "none"
    assert parse_page_range("1-3,7,9-", 10) == {1, 2, 3, 7, 9, 10}


def test_skip_heuristic_never_drops_priced_or_unreadable_pages():
    assert skip_reason_for("Table of Contents\nBottles 4\nMugs 9", 2, "native")
    assert skip_reason_for("Contact us for MRP ₹450 bottles", 2, "native") is None
    assert skip_reason_for("", 5, "none") is None
    assert skip_reason_for("Index\nModel SB-750 page 4", 3, "native") is None   # SKU present
